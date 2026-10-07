import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

# 1. SIMULACIÓN DE ESTRUCTURA CON SUBCARPETAS Y MÚLTIPLES HOJAS
base_dir = "reportes_anuales"

# Creamos estructura de carpetas: reportes_anuales/2026/Q1/
carpeta_q1 = os.path.join(base_dir, "2026", "Q1")
os.makedirs(carpeta_q1, exist_ok=True)

# Creamos un Excel con 2 pestañas (Enero y Febrero)
archivo_q1 = os.path.join(carpeta_q1, "ventas_Q1.xlsx")

with pd.ExcelWriter(archivo_q1, engine="openpyxl") as writer:
    df_enero = pd.DataFrame({
        "Mes": ["Enero", "Enero"],
        "Vendedor": ["Carlos", "Ana"],
        "Monto": [3200.0, 1800.0]
    })
    df_febrero = pd.DataFrame({
        "Mes": ["Febrero", "Febrero"],
        "Vendedor": ["Carlos", "Ana"],
        "Monto": [2900.0, 2100.0]
    })
    
    df_enero.to_excel(writer, sheet_name="Enero", index=False)
    df_febrero.to_excel(writer, sheet_name="Febrero", index=False)

# 2. LECTURA RECURSIVA CON OS.WALK Y PARSEO DE TODAS LAS HOJAS
registros_totales = []

# os.walk recorre la carpeta raíz y todas sus subcarpetas
for root, dirs, files in os.walk(base_dir):
    for file in files:
        if file.endswith(".xlsx"):
            ruta_completa = os.path.join(root, file)
            
            # pd.read_excel con sheet_name=None lee TODAS las pestañas
            hojas_excel = pd.read_excel(ruta_completa, sheet_name=None)
            
            for nombre_hoja, df_hoja in hojas_excel.items():
                registros_totales.append(df_hoja)

# Unir todas las pestañas de todos los archivos en un solo DataFrame
df_consolidado = pd.concat(registros_totales, ignore_index=True)

# Agrupar total de ventas por Mes
df_resumen = df_consolidado.groupby("Mes", as_index=False)["Monto"].sum()

# Guardar base limpia a Excel
nombre_archivo = "Consolidado_Anual_Avanzado.xlsx"
df_resumen.to_excel(nombre_archivo, index=False)

# 3. DISEÑO Y GRÁFICO CON OPENPYXL
wb = load_workbook(nombre_archivo)
ws = wb.active

fuente_encabezado = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
relleno_encabezado = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

for cell in ws[1]:
    cell.font = fuente_encabezado
    cell.fill = relleno_encabezado
    cell.alignment = Alignment(horizontal="center")

for row in range(2, ws.max_row + 1):
    ws.cell(row=row, column=2).number_format = '"$"#,##0.00'

for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

# Gráfico de barras de Ventas por Mes
chart = BarChart()
chart.type = "col"
chart.style = 10
chart.title = "Ventas Consolidadas por Mes"
chart.y_axis.title = "Monto Total ($)"
chart.x_axis.title = "Mes"

data = Reference(ws, min_col=2, min_row=1, max_row=ws.max_row)
categories = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)

chart.add_data(data, titles_from_data=True)
chart.set_categories(categories)
chart.legend = None

ws.add_chart(chart, "D2")

wb.save(nombre_archivo)
print("¡Consolidación avanzada exitosa! Procesadas subcarpetas y múltiples pestañas.")