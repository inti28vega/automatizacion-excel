import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

# 1. SIMULACIÓN Y LIMPIEZA DE DATOS CON PANDAS
datos_ventas = {
    "Vendedor": ["  carlos  ", "ana ", "MARIA", "  carlos  ", "ana "],
    "Producto": [" Laptops ", "Teclados", "monitores ", "Teclados", " Laptops "],
    "Monto": [1200.0, 150.0, None, 300.0, 2400.0],
    "Estado": ["Completado", "Completado", "Pendiente", "Completado", "Completado"]
}

df = pd.DataFrame(datos_ventas)

# Limpieza de texto: espacios extra fuera y formato Título (Capitalizado)
df["Vendedor"] = df["Vendedor"].str.strip().str.title()
df["Producto"] = df["Producto"].str.strip().str.title()

# Tratamiento de nulos: reemplazar montos vacíos por 0
df["Monto"] = df["Monto"].fillna(0)

# Agrupar total de ventas completadas por Vendedor
df_resumen = df[df["Estado"] == "Completado"].groupby("Vendedor", as_index=False)["Monto"].sum()

# Guardar base limpia a Excel
nombre_archivo = "Reporte_Ejecutivo_Ventas.xlsx"
df_resumen.to_excel(nombre_archivo, index=False)

# 2. DISEÑO CORPORATIVO CON OPENPYXL
wb = load_workbook(nombre_archivo)
ws = wb.active

# Estilos de encabezado
fuente_encabezado = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
relleno_encabezado = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

for cell in ws[1]:
    cell.font = fuente_encabezado
    cell.fill = relleno_encabezado
    cell.alignment = Alignment(horizontal="center")

# Formato de moneda para la columna Monto (Columna B)
for row in range(2, ws.max_row + 1):
    ws.cell(row=row, column=2).number_format = '"$"#,##0.00'

# Autoajustar ancho de columnas
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

# 3. GRÁFICO AUTOMÁTICO
chart = BarChart()
chart.type = "col"
chart.style = 10
chart.title = "Ventas Totales por Vendedor"
chart.y_axis.title = "Monto Total ($)"
chart.x_axis.title = "Vendedor"

data = Reference(ws, min_col=2, min_row=1, max_row=ws.max_row)
categories = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)

chart.add_data(data, titles_from_data=True)
chart.set_categories(categories)
chart.legend = None

# Insertar el gráfico al lado de la tabla de resumen
ws.add_chart(chart, "D2")

# Guardar versión final
wb.save(nombre_archivo)
print("¡Caso integrador ejecutado con éxito! Generado 'Reporte_Ejecutivo_Ventas.xlsx'.")