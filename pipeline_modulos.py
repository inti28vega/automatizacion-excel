import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def extraer_y_consolidar_datos(directorio_base):
    """Recorre carpetas anidadas y consolida todos los Excel/CSV encontrados."""
    registros = []
    for root, _, files in os.walk(directorio_base):
        for file in files:
            ruta_completa = os.path.join(root, file)
            if file.endswith(".xlsx") or file.endswith(".xls"):
                hojas = pd.read_excel(ruta_completa, sheet_name=None)
                for _, df in hojas.items():
                    registros.append(df)
            elif file.endswith(".csv"):
                df = pd.read_csv(ruta_completa)
                registros.append(df)
    
    if registros:
        return pd.concat(registros, ignore_index=True)
    return pd.DataFrame()

def auditar_y_limpiar_datos(df, col_vendedor, col_monto):
    """Limpia textos, aisla errores y retorna DataFrames usando columnas dinámicas."""
    df_temp = df.copy()
    
    if col_vendedor in df_temp.columns:
        df_temp[col_vendedor] = df_temp[col_vendedor].astype(str).str.strip()
    
    nulos = df_temp[df_temp[col_vendedor] == "None"] if col_vendedor in df_temp.columns else pd.DataFrame()
    montos_invalidos = df_temp[df_temp[col_monto] <= 0] if col_monto in df_temp.columns else pd.DataFrame()
    
    df_errores = pd.concat([nulos, montos_invalidos]).drop_duplicates()
    
    df_limpio = df_temp[
        (df_temp[col_vendedor] != "None") & 
        (df_temp[col_monto] > 0)
    ].drop_duplicates() if col_monto in df_temp.columns else df_temp
    
    return df_limpio, df_errores

def aplicar_formato_excel(archivo_salida, formato_moneda):
    """Aplica estilos profesionales, bordes negros definidos, formato de moneda y fila de Total General."""
    wb = load_workbook(archivo_salida)
    ws = wb.active

    # Definición de fuentes y rellenos
    fuente_encabezado = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    relleno_encabezado = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    
    # Bordes negros bien marcados
    borde_negro = Side(border_style="thin", color="000000")
    estilo_borde = Border(left=borde_negro, right=borde_negro, top=borde_negro, bottom=borde_negro)

    # 1. Estilo y bordes de encabezados
    for cell in ws[1]:
        cell.font = fuente_encabezado
        cell.fill = relleno_encabezado
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = estilo_borde

    # 2. Formato de moneda y bordes para las filas de datos
    for row in range(2, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            celda = ws.cell(row=row, column=col)
            celda.border = estilo_borde
            if col == 2:
                celda.number_format = formato_moneda

    # 3. Fila de Total General al final
    ultima_fila = ws.max_row + 1
    
    # Etiqueta "Total General"
    celda_etiqueta = ws.cell(row=ultima_fila, column=1, value="Total General")
    celda_etiqueta.font = Font(name="Calibri", size=11, bold=True)
    celda_etiqueta.border = estilo_borde
    
    # Fórmula de suma en Excel para la columna B
    celda_suma_rango = f"B2:B{ultima_fila - 1}"
    celda_total = ws.cell(row=ultima_fila, column=2, value=f"=SUM({celda_suma_rango})")
    celda_total.font = Font(name="Calibri", size=11, bold=True)
    celda_total.number_format = formato_moneda
    celda_total.border = estilo_borde

    # 4. Auto-ajuste del ancho de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    wb.save(archivo_salida)