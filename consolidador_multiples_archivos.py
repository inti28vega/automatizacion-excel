import os
import glob
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

# 1. CREAR SIMULACIÓN DE MÚLTIPLES ARCHIVOS EN UNA CARPETA
carpeta_sucursales = "sucursales_datos"
os.makedirs(carpeta_sucursales, exist_ok=True)

# Simulamos 3 archivos CSV de distintas sucursales
sucursal_1 = pd.DataFrame({
    "Sucursal": ["Sucursal Centro", "Sucursal Centro"],
    "Monto": [1500.0, 800.0],
    "Estado": ["Completado", "Completado"]
})

sucursal_2 = pd.DataFrame({
    "Sucursal": ["Sucursal Norte", "Sucursal Norte"],
    "Monto": [2300.0, 450.0],
    "Estado": ["Completado", "Pendiente"]
})

sucursal_3 = pd.DataFrame({
    "Sucursal": ["Sucursal Sur", "Sucursal Sur"],
    "Monto": [1100.0, 950.0],
    "Estado": ["Completado", "Completado"]
})

sucursal_1.to_csv(os.path.join(carpeta_sucursales, "ventas_centro.csv"), index=False)
sucursal_2.to_csv(os.path.join(carpeta_sucursales, "ventas_norte.csv"), index=False)
sucursal_3.to_csv(os.path.join(carpeta_sucursales, "ventas_sur.csv"), index=False)

# 2. LECTURA Y CONSOLIDACIÓN AUTOMÁTICA CON GLOB Y PANDAS
archivos_csv = glob.glob(os.path.join(carpeta_sucursales, "*.csv"))

lista_dataframes = []
for archivo in archivos_csv:
    df_temp = pd.read_csv(archivo)
    lista_dataframes.append(df_temp)

# Unir todos los DataFrames en uno solo
df_consolidado = pd.concat(lista_dataframes, ignore_index=True)

# Filtrar solo ventas completadas y agrupar por Sucursal
df_resumen = df_consolidado[df_consolidado["Estado"] == "Completado"].groupby("Sucursal", as_index=False)["Monto"].sum()

# Guardar base consolidada a Excel
nombre_archivo = "Consolidado_Sucursales.xlsx"
df_resumen.to_excel(nombre_archivo, index=False)

# 3. DISEÑO Y GRÁFICO CON OPENPYXL
wb = load_workbook(nombre_archivo)
ws = wb.active

# Estilos de encabezado
fuente_encabezado = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
relleno_encabezado = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

for cell in ws[1]:
    cell.font = fuente_encabezado
    cell.fill = relleno_encabezado
    cell.alignment = Alignment(horizontal="center")

# Formato de moneda para la columna Monto (Columna B)
for row in range(2, ws.max_row + 1):
    ws.cell(row=row, column=2).number_format = '"$"#,##0.00'

# Autoajustar ancho de columnas
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

# Gráfico automático
chart = BarChart()
chart.type = "col"
chart.style = 10
chart.title = "Ventas Completadas por Sucursal"
chart.y_axis.title = "Monto Total ($)"
chart.x_axis.title = "Sucursal"

data = Reference(ws, min_col=2, min_row=1, max_row=ws.max_row)
categories = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)

chart.add_data(data, titles_from_data=True)
chart.set_categories(categories)
chart.legend = None

ws.add_chart(chart, "D2")

wb.save(nombre_archivo)
print("¡Consolidación exitosa! Generado 'Consolidado_Sucursales.xlsx' a partir de múltiples archivos CSV.")