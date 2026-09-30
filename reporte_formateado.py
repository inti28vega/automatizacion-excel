import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 1. Crear y limpiar los datos con pandas
datos = {
    "Fecha": ["2026-09-01", "2026-09-02", "2026-09-02", "2026-09-04", "2026-09-05"],
    "Concepto": ["  papelería ", "Internet ", "Mantenimiento", "Alquiler", "Suministros  "],
    "Monto": [45.0, 120.0, 350.0, 500.0, 85.5],
    "Estado": ["Aprobado", "Aprobado", "Pendiente", "Aprobado", "Pendiente"]
}

df = pd.DataFrame(datos)
df["Concepto"] = df["Concepto"].str.strip().str.capitalize()

# Guardar temporalmente a Excel
nombre_archivo = "Reporte_Estilizado.xlsx"
df.to_excel(nombre_archivo, index=False)

# 2. Cargar el libro con openpyxl para darle diseño
wb = load_workbook(nombre_archivo)
ws = wb.active

# Estilos visuales
fuente_encabezado = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
relleno_encabezado = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Azul corporativo
alineacion_centro = Alignment(horizontal="center", vertical="center")

# Aplicar estilo a la fila de encabezados (Fila 1)
for cell in ws[1]:
    cell.font = fuente_encabezado
    cell.fill = relleno_encabezado
    cell.alignment = alineacion_centro

# Aplicar formato de moneda a la columna 'Monto' (Columna C / columna 3)
for row in range(2, ws.max_row + 1):
    celda_monto = ws.cell(row=row, column=3)
    celda_monto.number_format = '"$"#,##0.00'

# Autoajustar el ancho de las columnas según el contenido
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# 3. Guardar el archivo final formateado
wb.save(nombre_archivo)
print("¡Reporte estilizado y guardado con éxito como 'Reporte_Estilizado.xlsx'!")