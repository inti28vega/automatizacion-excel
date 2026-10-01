import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

# 1. Crear y limpiar los datos con pandas
datos = {
    "Concepto": ["Papelería", "Internet", "Mantenimiento", "Alquiler", "Suministros"],
    "Monto": [45.0, 120.0, 350.0, 500.0, 85.5]
}

df = pd.DataFrame(datos)

# Guardar en Excel preliminar
nombre_archivo = "Reporte_Con_Grafico.xlsx"
df.to_excel(nombre_archivo, index=False)

# 2. Cargar el libro con openpyxl para diseño y gráfico
wb = load_workbook(nombre_archivo)
ws = wb.active

# Estilos visuales a los encabezados
fuente_encabezado = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
relleno_encabezado = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

for cell in ws[1]:
    cell.font = fuente_encabezado
    cell.fill = relleno_encabezado
    cell.alignment = Alignment(horizontal="center")

# Formato de moneda en la columna Monto
for row in range(2, ws.max_row + 1):
    ws.cell(row=row, column=2).number_format = '"$"#,##0.00'

# Autoajustar ancho de columnas
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# 3. CREACIÓN DEL GRÁFICO DE BARRAS
chart = BarChart()
chart.type = "col"  # Gráfico de columnas verticales
chart.style = 10   # Estilo de diseño prediseñado de Excel
chart.title = "Gastos Totales por Concepto"
chart.y_axis.title = "Monto ($)"
chart.x_axis.title = "Concepto"

# Referencia de los datos numéricos (Columna 2: Monto, desde fila 1 hasta la última)
data = Reference(ws, min_col=2, min_row=1, max_row=ws.max_row)

# Referencia de las etiquetas/categorías (Columna 1: Concepto, desde fila 2 hasta la última)
categories = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)

# Vincular datos y categorías al gráfico
chart.add_data(data, titles_from_data=True)
chart.set_categories(categories)

# Quitar la leyenda lateral (no es necesaria para una sola serie de datos)
chart.legend = None

# Insertar el gráfico en la celda D2 (al lado de la tabla)
ws.add_chart(chart, "D2")

# 4. Guardar el archivo final
wb.save(nombre_archivo)
print("¡Reporte con gráfico generado con éxito en 'Reporte_Con_Grafico.xlsx'!")